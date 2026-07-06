from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile


SUPPORTED_DOCUMENT_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx"}
HEADER_PATTERN = re.compile(r"(?m)^(#{1,6})\s+(.+?)\s*$")
ARTICLE_PATTERN = re.compile(r"(?m)^第[一二三四五六七八九十百千零〇\d]+条\s+")


class DocumentParseError(RuntimeError):
    pass


@dataclass(frozen=True)
class SourceDocument:
    source_path: str
    title: str
    text: str
    metadata: dict[str, object] = field(default_factory=dict)


@dataclass(frozen=True)
class DocumentChunk:
    chunk_id: str
    title: str
    content: str
    metadata: dict[str, object] = field(default_factory=dict)


def load_policy_document_chunks(path: Path) -> list[DocumentChunk]:
    documents = load_source_documents(path)
    chunks: list[DocumentChunk] = []
    next_id = 1
    for document in documents:
        for chunk in split_policy_document(document):
            chunks.append(
                DocumentChunk(
                    chunk_id=f"{document.source_path}:{next_id}",
                    title=chunk.title,
                    content=chunk.content,
                    metadata=chunk.metadata,
                )
            )
            next_id += 1
    return chunks


def load_source_documents(path: Path) -> list[SourceDocument]:
    if not path.exists():
        return []
    if path.is_file():
        return [_load_source_document(path)]

    documents: list[SourceDocument] = []
    for candidate in sorted(item for item in path.rglob("*") if item.is_file()):
        if candidate.suffix.lower() not in SUPPORTED_DOCUMENT_SUFFIXES:
            continue
        documents.append(_load_source_document(candidate))
    return documents


def split_policy_document(document: SourceDocument) -> list[DocumentChunk]:
    text = normalize_text(document.text)
    if not text:
        return []

    header_matches = list(HEADER_PATTERN.finditer(text))
    chunks: list[DocumentChunk] = []
    if not header_matches:
        return _split_article_chunks(document.title, text, document)

    prefix = text[: header_matches[0].start()].strip()
    if prefix:
        chunks.extend(_split_article_chunks(document.title, prefix, document))

    heading_stack: list[tuple[int, str]] = []
    for index, match in enumerate(header_matches):
        level = len(match.group(1))
        heading = match.group(2).strip()
        heading_stack = [(item_level, title) for item_level, title in heading_stack if item_level < level]
        heading_stack.append((level, heading))

        content_start = match.end()
        content_end = header_matches[index + 1].start() if index + 1 < len(header_matches) else len(text)
        content = text[content_start:content_end].strip()
        if not content:
            continue
        section_title = " / ".join(title for _, title in heading_stack)
        chunks.extend(_split_article_chunks(section_title, content, document))
    return chunks


def normalize_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = re.sub(r"[ \t]+\n", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _load_source_document(path: Path) -> SourceDocument:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return SourceDocument(
            source_path=str(path),
            title=path.stem,
            text=path.read_text(encoding="utf-8"),
            metadata={"format": suffix.lstrip(".")},
        )
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    raise DocumentParseError(f"unsupported policy document format: {path}")


def _load_docx(path: Path) -> SourceDocument:
    try:
        from docx import Document as DocxDocument  # type: ignore

        document = DocxDocument(str(path))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        text = "\n".join(blocks)
        parser = "python-docx"
    except ImportError:
        text = _load_docx_with_stdlib(path)
        parser = "zip-xml"
    return SourceDocument(
        source_path=str(path),
        title=path.stem,
        text=text,
        metadata={"format": "docx", "parser": parser},
    )


def _load_docx_with_stdlib(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            raw_xml = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise DocumentParseError(f"failed to parse docx document: {path}") from exc

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ET.fromstring(raw_xml)
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", namespace):
        texts = [node.text or "" for node in paragraph.findall(".//w:t", namespace)]
        paragraph_text = "".join(texts).strip()
        if paragraph_text:
            paragraphs.append(paragraph_text)
    return "\n".join(paragraphs)


def _load_pdf(path: Path) -> SourceDocument:
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise DocumentParseError("PDF parsing requires PyMuPDF; install python_agent/requirements-rag.txt") from exc

    pages: list[str] = []
    with fitz.open(path) as document:  # type: ignore[attr-defined]
        for page_index, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            if page_text:
                pages.append(f"[page {page_index}]\n{page_text}")

    return SourceDocument(
        source_path=str(path),
        title=path.stem,
        text="\n\n".join(pages),
        metadata={
            "format": "pdf",
            "parser": "pymupdf",
            "ocr_required": not bool(pages),
        },
    )


def _split_article_chunks(title: str, content: str, document: SourceDocument) -> list[DocumentChunk]:
    article_matches = list(ARTICLE_PATTERN.finditer(content))
    if not article_matches:
        return [
            DocumentChunk(
                chunk_id="",
                title=title,
                content=content,
                metadata=dict(document.metadata, source_path=document.source_path),
            )
        ]

    chunks: list[DocumentChunk] = []
    for index, match in enumerate(article_matches):
        start = match.start()
        end = article_matches[index + 1].start() if index + 1 < len(article_matches) else len(content)
        article = content[start:end].strip()
        first_line = article.splitlines()[0].strip() if article else ""
        article_title = first_line[:48] if first_line else f"条款 {index + 1}"
        chunks.append(
            DocumentChunk(
                chunk_id="",
                title=f"{title} / {article_title}",
                content=article,
                metadata=dict(document.metadata, source_path=document.source_path, article_title=article_title),
            )
        )
    return chunks
